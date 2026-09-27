"""Document Text Extractor for Plain Text, DOCX, and PDF Files."""
import io
import os
from typing import Tuple


def extract_text_from_file(filename: str, file_bytes: bytes) -> str:
    """Extract plain text from .txt, .docx, or .pdf files.
    
    Args:
        filename: Name of the file with extension.
        file_bytes: Raw binary content of the file.
        
    Returns:
        Extracted text as a unicode string.
    """
    _, ext = os.path.splitext(filename.lower())
    
    if ext == ".docx":
        try:
            import docx
            doc = docx.Document(io.BytesIO(file_bytes))
            paragraphs = []
            for p in doc.paragraphs:
                paragraphs.append(p.text)
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    if row_text:
                        paragraphs.append(row_text)
            return "\n".join(paragraphs)
        except Exception as e:
            raise ValueError(f"Failed to parse DOCX file: {str(e)}")
            
    elif ext == ".pdf":
        try:
            import pypdf
            reader = pypdf.PdfReader(io.BytesIO(file_bytes))
            pages_text = []
            for page in reader.pages:
                extracted = page.extract_text()
                if extracted:
                    pages_text.append(extracted)
            return "\n\n".join(pages_text)
        except Exception as e:
            raise ValueError(f"Failed to parse PDF file: {str(e)}")
            
    else:
        # Default text file decoding (UTF-8, UTF-16, Latin-1)
        for encoding in ("utf-8", "utf-16", "latin-1", "ascii"):
            try:
                return file_bytes.decode(encoding)
            except (UnicodeDecodeError, AttributeError):
                continue
        return file_bytes.decode("utf-8", errors="replace")
